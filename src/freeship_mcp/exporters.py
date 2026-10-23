"""Portable report writers consuming only native engine results."""
from __future__ import annotations
import csv
import json
from pathlib import Path
from .linesplan import linesplan_dxf
from openpyxl import Workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FIELDS = [
    "draft_m", "displacement_t", "volume_geometry_m3", "volume_corrected_m3",
    "lcb_ap_m", "vcb_m", "lcf_ap_m", "bm_transverse_m", "bm_longitudinal_m",
    "waterline_length_m", "waterline_beam_m", "waterplane_area_m2",
    "block_coefficient", "waterplane_coefficient", "prismatic_coefficient",
    "midship_coefficient", "wetted_surface_m2", "trim_m", "heel_degrees", "no_immersion",
    "block_coefficient_project_geometry", "waterplane_coefficient_project",
    "minimum_open_boundary_z_m", "minimum_freeboard_m",
]
CHART_FIELDS = ["displacement_t", "lcb_ap_m", "vcb_m", "lcf_ap_m",
                "bm_transverse_m", "bm_longitudinal_m", "waterline_length_m", "waterplane_area_m2"]


def hydrostatics_report(path: Path, rows: list[dict], display_decimals: int = 2, provenance: dict | None = None) -> dict:
    if not rows:
        raise ValueError("At least one hydrostatics row is required")
    if not 0 <= display_decimals <= 8:
        raise ValueError("display_decimals must be between 0 and 8")
    if path.suffix.lower() == ".csv":
        with path.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.DictWriter(stream, fieldnames=FIELDS, extrasaction="ignore")
            writer.writeheader()
            writer.writerows(rows)
        return {"rows": len(rows), "charts": 0}
    book = Workbook()
    data = book.active
    data.title = "Hydrostatics"
    data.append(FIELDS)
    for row in rows:
        data.append([row.get(field) for field in FIELDS])
    edge = Side(style="thin", color="000000")
    for cell in data[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="24465A")
    for row in data:
        for cell in row:
            cell.border = Border(left=edge, right=edge, top=edge, bottom=edge)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            if cell.row > 1 and cell.data_type == "n":
                cell.number_format = "0" + ("." + "0" * display_decimals if display_decimals else "")
    for col in range(1, len(FIELDS) + 1):
        data.column_dimensions[get_column_letter(col)].width = 22
    data.row_dimensions[1].height = 45
    data.freeze_panes = "B2"
    data.auto_filter.ref = data.dimensions
    data.print_title_rows = "1:1"
    data.sheet_properties.pageSetUpPr.fitToPage = True
    data.page_setup.orientation = "landscape"
    data.page_setup.paperSize = data.PAPERSIZE_A3
    data.page_setup.fitToWidth = 1
    data.page_setup.fitToHeight = 0
    plots = book.create_sheet("Curves")
    for i, field in enumerate(CHART_FIELDS):
        chart = ScatterChart()
        chart.title = field
        chart.x_axis.title = "Draft, m"
        chart.y_axis.title = field
        chart.scatterStyle = "lineMarker"
        x = Reference(data, min_col=1, min_row=2, max_row=len(rows) + 1)
        y = Reference(data, min_col=FIELDS.index(field) + 1, min_row=2, max_row=len(rows) + 1)
        chart.series.append(Series(y, x, title=field))
        chart.width, chart.height = 19, 10
        plots.add_chart(chart, f"{'A' if i % 2 == 0 else 'L'}{1 + (i // 2) * 21}")
    metadata = book.create_sheet("Provenance")
    metadata.append(["field", "JSON value"])
    for key, value in (provenance or {}).items():
        encoded = json.dumps(value, ensure_ascii=False, allow_nan=False)
        for start in range(0, len(encoded), 16000):
            metadata.append([key if start == 0 else key+" (continued)", encoded[start:start+16000]])
    metadata.append(["coefficient_basis", json.dumps(rows[0].get("coefficient_basis", {}), ensure_ascii=False)])
    metadata.column_dimensions["A"].width = 25
    metadata.column_dimensions["B"].width = 100
    metadata.freeze_panes = "A2"
    book.save(path)
    return {"rows": len(rows), "charts": len(CHART_FIELDS)}
