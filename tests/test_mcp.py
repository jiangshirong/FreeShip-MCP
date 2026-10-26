import asyncio
import json
from pathlib import Path
import shutil
import sys
import os
import pytest
import ezdxf
from openpyxl import load_workbook
from mcp import ClientSession
from mcp.client.stdio import StdioServerParameters, stdio_client
from freeship_mcp.exporters import FIELDS

ROOT = Path(__file__).resolve().parents[1]


@pytest.mark.parametrize("runtime", ["python", "standalone"])
def test_real_stdio_full_chain(tmp_path, runtime):
    shutil.copy2(ROOT / "examples/synthetic-hull.fbm", tmp_path / "hull.fbm")
    async def workflow():
        if runtime == "python":
            command, args, env = sys.executable, ["-m", "freeship_mcp.server"], None
        else:
            binary = ROOT / "dist/standalone/freeship-mcp/freeship-mcp.exe"
            if not binary.exists():
                pytest.skip("Build the standalone package first")
            command, args = str(binary), []
            # Server/native engine cannot find Python, Pascal or desktop FreeShip on PATH.
            env = {"PATH": str(Path(os.environ["SystemRoot"]) / "System32"), "PYTHONPATH": "", "PYTHONHOME": ""}
        params = StdioServerParameters(command=command, args=args + ["--workspace", str(tmp_path)], env=env)
        async with stdio_client(params) as (reader, writer):
            async with ClientSession(reader, writer) as client:
                await client.initialize()
                tools = await client.list_tools()
                assert len(tools.tools) == 15
                model_tools={'save_model','scale_model','edit_control_net','generate_sections','rebuild_hull'}
                for tool in json.loads(tools.model_dump_json(by_alias=True))['tools']:
                    if tool['name'] in model_tools:
                        assert tool['inputSchema']['properties']['output_version']['default']==335
                        assert 'output_path' in tool['inputSchema']['required']
                async def call(name, **kwargs):
                    result = await client.call_tool(name, kwargs)
                    assert not result.is_error, result
                    if result.structured_content is not None:
                        return result.structured_content
                    return json.loads(result.content[0].text)
                info = await call("engine_info")
                assert info["headless"]
                net=await call("inspect_model", model_path="hull.fbm", include_control_net=True)
                index=net['control_net_points'][0]['index']
                x=net['control_net_points'][0]['coordinate_m'][0]
                local=await call("edit_control_net",model_path="hull.fbm",output_path="local.fbm",
                    expected_model_sha256=net['provenance']['model']['sha256'],
                    edit={'action':'move_points','points':[{'index':index,'x_m':x-.01}]})
                assert local['changed_points'][0]['after_m'][0]==pytest.approx(x-.01,abs=1e-6)
                stale=await client.call_tool('edit_control_net',{'model_path':'local.fbm','output_path':'stale.fbm',
                    'expected_model_sha256':net['provenance']['model']['sha256'],
                    'edit':{'action':'translate_points','point_indices':[index],'dx_m':1}})
                assert stale.is_error and not (tmp_path/'stale.fbm').exists()
                quality=await call('analyze_hull_quality',model_path='local.fbm',stations_m=[10,20,30],
                    waterlines_m=[1,2,3],buttocks_m=[1,2],max_turn_degrees=20,
                    alignment_rules=[{'point_indices':[index],'axis':'x','value_m':x-.01,'tolerance_m':1e-5}])
                assert quality['curve_count']>0 and quality['alignment_checks'][0]['status']=='passed'
                partial=await call('analyze_hull_quality',model_path='local.fbm',stations_m=[1000],
                    waterlines_m=[1],max_turn_degrees=180)
                assert partial['status']=='incomplete' and partial['curve_checks']=='partial'
                await call("validate_model", model_path="hull.fbm")
                saved=await call("save_model", model_path="hull.fbm", output_path="saved.ftm")
                assert saved['output_format_version']=='fv335' and Path(saved['output_path']).name=='saved.ftm'
                override=await call('save_model',model_path='hull.fbm',output_path='explicit.fbm',output_version=462)
                assert override['output_format_version']=='fv462'
                await call("scale_model", model_path="saved.ftm", output_path="scaled.fbm", x_factor=1.1, y_factor=1.1, z_factor=1.1)
                h = await call("calculate_hydrostatics", model_path="scaled.fbm", draft_m=3.3)
                check = await call("verify_model_targets", model_path="scaled.fbm", draft_m=3.3,
                    targets=[{"metric":"project_length_m", "expected":44, "absolute_tolerance":.001},
                             {"metric":"displacement_t", "expected":h["displacement_t"], "absolute_tolerance":.001},
                             {"metric":"section_top_above_keel_m", "section_x_m":22, "expected":6.6, "absolute_tolerance":.01}])
                assert check["status"] == "passed", check
                assert check["checks"][0]["basis"]["geometry_verified"] is False
                await call("rebuild_hull", model_path="scaled.fbm", output_path="rebuilt.fbm",
                           displacement_t=h["displacement_t"] * 1.01, lcb_value=h["lcb_ap_m"],
                           lcb_reference="aft_perpendicular_m", layer_indices=[0])
                section = await call("generate_sections", model_path="rebuilt.fbm", output_path="sections.fbm",
                    stations_m=[44 * i / 10 for i in range(11)],
                    waterlines_m=[0.66 * i for i in range(1, 7)], buttocks_m=[1, 2, 3])
                assert section["stations"] == 11
                ranged = await call("calculate_hydrostatics_series", model_path="sections.fbm", draft_range_m=[0, 3, 1])
                assert [row["draft_m"] for row in ranged["rows"]] == [0,1,2,3]
                dx = await call("export_linesplan_dxf", model_path="sections.fbm", output_path="lines.dxf",
                    options={"title":"Synthetic hull","body_split_x_m":22,"include_grid":True})
                assert dx["views"]["BODY"] > 0
                assert dx['boundary_entities']['PROFILE']>0 and dx['grid_entities']['BODY']>0
                assert dx['nonempty_sections']['station']>=9
                series = await call("calculate_hydrostatics_series", model_path="sections.fbm", drafts_m=[0, 1, 2, 3.3])
                await call("export_hydrostatics", model_path="sections.fbm", output_path="hydro.xlsx", drafts_m=[0, 1, 2, 3.3])
                await call("export_hydrostatics", model_path="sections.fbm", output_path="hydro.csv", drafts_m=[0, 1, 2, 3.3])
                for exported in ("hydro.xlsx", "hydro.csv", "lines.dxf"):
                    check = await call("check_export_source", model_path="sections.fbm", export_path=exported)
                    assert check["status"] == "current"
                wrong = await call("check_export_source", model_path="hull.fbm", export_path="hydro.xlsx")
                assert wrong["status"] == "stale_or_modified" and not wrong["model_matches"]
                with (tmp_path/"hydro.csv").open("a",encoding="utf-8") as stream:stream.write("tampered")
                wrong = await call("check_export_source", model_path="sections.fbm", export_path="hydro.csv")
                assert wrong["status"] == "stale_or_modified" and not wrong["artifact_matches"]
                bad = await client.call_tool("inspect_model", {"model_path": "../outside.fbm"})
                assert bad.is_error
                return series["rows"]
    rows = asyncio.run(workflow())
    book = load_workbook(tmp_path / "hydro.xlsx", data_only=True)
    assert len(book["Curves"]._charts) == 8
    for i, row in enumerate(rows, 2):
        for j, field in enumerate(FIELDS, 1):
            actual = book["Hydrostatics"].cell(i, j).value
            expected = row.get(field)
            if isinstance(expected, bool):
                assert actual == expected
            elif expected is not None:
                assert actual == pytest.approx(expected, rel=1e-14, abs=1e-14)
            else:
                assert actual is None
    assert book['Provenance']['A2'].value == 'schema_version'
    for chart in book["Curves"]._charts:
        assert chart.series[0].xVal.numRef.f == "'Hydrostatics'!$A$2:$A$5"
    assert book["Hydrostatics"].print_title_rows == "$1:$1"
    doc = ezdxf.readfile(tmp_path / "lines.dxf")
    assert not doc.audit().has_errors
    assert doc.units == ezdxf.units.M
    for entity in doc.modelspace().query("LWPOLYLINE"):
        layer = entity.dxf.layer
        if layer.startswith("PROFILE_STATION_"):
            index = int(layer.rsplit("_", 1)[1])
            for vertex in entity.get_points("xy"):
                assert vertex[0] == pytest.approx(44 * index / 10, abs=1e-4)
        if layer.startswith("PROFILE_"):
            for x, y in entity.get_points("xy"):
                assert -0.1 <= x <= 44.1
                assert -0.1 <= y <= 6.7
